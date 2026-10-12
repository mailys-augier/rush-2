"""Régénère les deux classeurs à partir des quatre exports de la pharmacie.

    output/Pharma_Analyse.xlsx     classeur du client : synthèse, ventes, calendrier, conseils, aide, données
    output/Pharma_Traitement.xlsx  classeur interne : détail des contrôles et du traitement des données

Usage (depuis la racine du dépôt) :
    python src/build_workbook.py
"""
import argparse
from math import ceil
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

ROOT = Path(__file__).resolve().parent.parent

# Nom court affiché au pharmacien pour chaque code ATC de l'export
GROUPES = {
    "M01AB": "Anti-inflammatoires type diclofénac",
    "M01AE": "Anti-inflammatoires type ibuprofène",
    "N02BA": "Aspirine et dérivés",
    "N02BE": "Paracétamol",
    "N05B": "Anxiolytiques",
    "N05C": "Somnifères et sédatifs",
    "R03": "Asthme et BPCO",
    "R06": "Antihistaminiques (allergies)",
}
G = list(GROUPES)

# Classification ATC de l'OMS (index ATC/DDD, https://atcddd.fhi.no/atc_ddd_index/), consultée le 09/10/2026.
# Colonnes : nom officiel OMS, traduction française, famille (niveaux supérieurs), exemples de substances.
ATC = {
    "M01AB": ("Acetic acid derivatives and related substances", "Dérivés de l'acide acétique et substances apparentées",
              "Système musculo-squelettique > Anti-inflammatoires et antirhumatismaux non stéroïdiens",
              "diclofénac, indométacine, kétorolac, acéclofénac, étodolac"),
    "M01AE": ("Propionic acid derivatives", "Dérivés de l'acide propionique",
              "Système musculo-squelettique > Anti-inflammatoires et antirhumatismaux non stéroïdiens",
              "ibuprofène, naproxène, kétoprofène, flurbiprofène, dexkétoprofène"),
    "N02BA": ("Salicylic acid and derivatives", "Acide salicylique et dérivés",
              "Système nerveux > Analgésiques > Autres analgésiques et antipyrétiques",
              "acide acétylsalicylique (aspirine), diflunisal, carbasalate calcique"),
    "N02BE": ("Anilides", "Anilides",
              "Système nerveux > Analgésiques > Autres analgésiques et antipyrétiques",
              "paracétamol, propacétamol, paracétamol en association"),
    "N05B": ("Anxiolytics", "Anxiolytiques", "Système nerveux > Psycholeptiques",
             "sous-groupes : dérivés des benzodiazépines (N05BA), dérivés du diphénylméthane (N05BB), carbamates "
             "(N05BC), autres anxiolytiques (N05BX)"),
    "N05C": ("Hypnotics and sedatives", "Hypnotiques et sédatifs", "Système nerveux > Psycholeptiques",
             "sous-groupes : dérivés des benzodiazépines (N05CD), médicaments apparentés aux benzodiazépines "
             "(N05CF), agonistes des récepteurs de la mélatonine (N05CH), barbituriques (N05CA)"),
    "R03": ("Drugs for obstructive airway diseases", "Médicaments des syndromes obstructifs des voies aériennes",
            "Système respiratoire",
            "sous-groupes : adrénergiques inhalés (R03A), autres médicaments inhalés (R03B), adrénergiques à usage "
            "systémique (R03C), autres médicaments systémiques (R03D)"),
    "R06": ("Antihistamines for systemic use", "Antihistaminiques à usage systémique", "Système respiratoire",
            "sous-groupe unique R06A, subdivisé selon la structure chimique"),
}

EXPORTS = {
    "Hourly": "%m/%d/%Y %H:%M",
    "Daily": "%m/%d/%Y",
    "Weekly": "%m/%d/%Y",
    "Monthly": "%Y-%m-%d",
}
JOURS = ["Lundi", "Mardi", "Mercredi", "Jeudi", "Vendredi", "Samedi", "Dimanche"]
MOIS = ["Janvier", "Février", "Mars", "Avril", "Mai", "Juin", "Juillet", "Août",
        "Septembre", "Octobre", "Novembre", "Décembre"]
TOL = 0.01  # écart toléré entre deux exports (arrondis)
SEUIL_FORT, SEUIL_CREUX = 1.15, 0.85  # indice saisonnier à partir duquel un mois est dit fort ou creux

# Onglets du classeur client dans l'ordre d'usage du pharmacien : (description, couleur claire, couleur foncée).
# Chaque onglet a sa couleur : fond clair et texte foncé. Tout bouton qui y mène et le titre de sa page la reprennent.
ONGLETS = {
    "Synthèse": ("Page d'accueil : conclusions, chiffres qui les prouvent, traitement des données", "E4DAF5", "5B3FA6"),
    "Ventes": ("Choisir des groupes et une période, voir les ventes correspondantes", "B9E6E0", "14756A"),
    "Calendrier": ("Mois forts et mois creux de chaque groupe", "F9C9DA", "B03A6B"),
    "Conseils": ("Pour chaque groupe : évolution, saison et conseil de commande", "C8EBD9", "1E7A52"),
    "Questions": ("Réponses aux questions les plus fréquentes", "FFD9B3", "B85C00"),
    "Médicaments": ("Ce que contient chaque groupe, selon la classification de l'OMS", "FFE9A8", "8A6A00"),
    "Méthode": ("Comment chaque chiffre est obtenu, et comment le refaire à la main", "F5D0C5", "A8452B"),
    "Statistiques": ("Tous les chiffres détaillés : par groupe, mois, jour, heure, année", "D5DCF7", "3C4FA8"),
    "État des données": ("Quel export sert à quoi, ce qui a été vérifié, corrigé ou écarté", "E3E3C3", "6B6B1F"),
    "Données": ("Ventes jour par jour, telles qu'exportées (lues par l'onglet Ventes)", "E0E0E0", "555555"),
}


def charger(dossier):
    exports = {}
    for nom, fmt in EXPORTS.items():
        df = pd.read_csv(dossier / f"Pharma_Ventes_{nom}.csv")
        df["datum"] = pd.to_datetime(df["datum"], format=fmt)
        exports[nom] = df
    return exports


def mois_complets(jour):
    """Mois dont tous les jours du calendrier sont présents dans l'export journalier."""
    nb = jour[G[0]].resample("ME").count()
    return nb.index[nb.values == nb.index.days_in_month]


def controler(exports):
    """Contrôle chaque export et les confronte entre eux. Renvoie (contrôles, détail mois par mois)."""
    h, d, w, m = (exports[n].set_index("datum") for n in EXPORTS)
    lignes = []

    def ajouter(controle, resultat, ok):
        lignes.append({"Contrôle": controle, "Résultat": resultat, "Statut": "OK" if ok else "À noter"})

    for nom, df in exports.items():
        nb_na = int(df.isna().sum().sum())
        nb_dup = int(df["datum"].duplicated().sum())
        nb_neg = int((df[G] < 0).sum().sum())
        ajouter(f"{nom} : valeurs manquantes / dates en double / quantités négatives",
                f"{nb_na} / {nb_dup} / {nb_neg}", nb_na + nb_dup + nb_neg == 0)

    jours_absents = pd.date_range(d.index.min(), d.index.max(), freq="D").difference(d.index)
    ajouter("Daily : jours absents entre le premier et le dernier jour", str(len(jours_absents)),
            len(jours_absents) == 0)
    heures_absentes = pd.date_range(h.index.min(), h.index.max(), freq="h").difference(h.index)
    ajouter("Hourly : heures absentes entre la première et la dernière heure", str(len(heures_absentes)),
            len(heures_absentes) == 0)

    def ecarts(fin, grossier):
        return grossier[G].sub(fin, fill_value=0)

    def en_ecart(e):
        return (e.abs() > TOL).any(axis=1)

    e_hd = ecarts(h[G].resample("D").sum(), d)
    ajouter("Hourly regroupé par jour = Daily", f"{int(en_ecart(e_hd).sum())} jour(s) en écart sur {len(d)}",
            not en_ecart(e_hd).any())
    e_dw = ecarts(d[G].resample("W-SUN").sum(), w)
    ajouter("Daily regroupé par semaine = Weekly", f"{int(en_ecart(e_dw).sum())} semaine(s) en écart sur {len(w)}",
            not en_ecart(e_dw).any())
    e_dm = ecarts(d[G].resample("ME").sum(), m)
    ajouter("Daily regroupé par mois = Monthly", f"{int(en_ecart(e_dm).sum())} mois en écart sur {len(m)}",
            not en_ecart(e_dm).any())

    jours_fichier = exports["Daily"]["Weekday Name"].values
    faux = int((d.index.day_name().values != jours_fichier).sum())
    ajouter("Daily : colonne « Weekday Name » cohérente avec la date", f"{faux} ligne(s) incohérente(s)", faux == 0)
    heures_daily = exports["Daily"]["Hour"]
    ajouter("Daily : colonne « Hour »",
            f"vaut {heures_daily.mode().iloc[0]} sur {int((heures_daily == heures_daily.mode().iloc[0]).sum())} "
            f"lignes sur {len(heures_daily)} : somme des heures, sans signification", False)

    complets = mois_complets(d)
    ajouter("Dernière vente enregistrée", h.index.max().strftime("%d/%m/%Y %H:%M"), True)
    ajouter("Dernier mois complet", complets.max().strftime("%m/%Y"), True)
    incomplets = sorted(set(d.index.to_period("M")) - set(complets.to_period("M")))
    ajouter("Mois incomplets dans le Daily", ", ".join(p.strftime("%m/%Y") for p in incomplets) or "aucun",
            not incomplets)
    entiers = bool((d[G] % 1 == 0).all().all())
    ajouter("Quantités entières", "oui" if entiers else "non : l'unité n'est pas la boîte", entiers)

    nb = d[G[0]].resample("ME").count().reindex(e_dm.index, fill_value=0)
    mois = pd.DataFrame({
        "Mois": e_dm.index.strftime("%m/%Y"),
        "Jours présents": nb.values,
        "Jours du mois": e_dm.index.days_in_month,
        "Complet": ["oui" if a == b else "non" for a, b in zip(nb.values, e_dm.index.days_in_month)],
        "Monthly en écart": ["oui" if x else "non" for x in en_ecart(e_dm)],
    })
    for g in G:
        mois[f"Écart {g} (Monthly - Daily)"] = e_dm[g].round(2).values
    return pd.DataFrame(lignes), mois


def etat_des_exports(exports):
    roles = {
        "Hourly": "Profil des ventes par heure uniquement",
        "Daily": "RÉFÉRENCE : toutes les statistiques, l'onglet Ventes, les mois et semaines recalculés",
        "Weekly": "Contrôle de cohérence uniquement (identique au Daily regroupé)",
        "Monthly": "ÉCARTÉ : ne concorde pas avec les trois autres exports",
    }
    return pd.DataFrame([{
        "Export": f"Pharma_Ventes_{nom}.csv",
        "Lignes": len(df),
        "Début": df["datum"].min().strftime("%d/%m/%Y"),
        "Fin": df["datum"].max().strftime("%d/%m/%Y"),
        "Utilisation": roles[nom],
    } for nom, df in exports.items()])


def decisions():
    return pd.DataFrame({"Corrigé ou écarté": [
        "Export Monthly écarté : les mois sont recalculés à partir du Daily.",
        "Mois incomplets exclus de toutes les statistiques mensuelles et de la saisonnalité.",
        "Dernière ligne du Weekly (semaine incomplète) non utilisée.",
        "Colonne « Hour » du Daily ignorée ; année, mois et jour de semaine recalculés à partir de la date.",
        "Aucune valeur de vente modifiée : aucune ligne supprimée, aucun plafonnement.",
        "Unité des quantités non documentée par le logiciel de caisse : les groupes ne sont pas "
        "additionnés ni comparés en volume entre eux.",
    ]})


def groupes_atc():
    df = pd.DataFrame.from_dict(ATC, orient="index", columns=[
        "Nom officiel OMS (anglais)", "Traduction française", "Famille", "Exemples ou sous-groupes"])
    df.insert(0, "Nom utilisé dans ce classeur", pd.Series(GROUPES))
    df["Source"] = [f"https://atcddd.fhi.no/atc_ddd_index/?code={g}" for g in df.index]
    return df.rename_axis("Code ATC").reset_index()


def methode():
    lignes = [
        ("1. Lire les exports",
         "Ouvre les 4 fichiers CSV et transforme la colonne « datum » en dates. Format mois/jour/année, sauf le "
         "Monthly en année-mois-jour.",
         "Données > À partir d'un fichier texte/CSV. Pour la colonne datum, choisir le type « Date » avec les "
         "paramètres régionaux Anglais (États-Unis), sinon le 2 janvier devient le 1er février."),
        ("2. Contrôler chaque export",
         "Compte les cases vides, les dates en double et les quantités négatives. Vérifie qu'aucun jour ni aucune "
         "heure ne manque.",
         "NB.VIDE sur la plage ; NB.SI(plage;\"<0\") ; pour les doublons, Mise en forme conditionnelle > Valeurs en "
         "double sur la colonne des dates ; pour les jours manquants, comparer (dernière date - première date + 1) "
         "au nombre de lignes."),
        ("3. Comparer les exports entre eux",
         "Additionne les heures par jour, les jours par semaine et les jours par mois, puis compare au fichier "
         "correspondant. Un écart inférieur à 0,01 est un arrondi.",
         "SOMME.SI.ENS sur le Daily entre le premier et le dernier jour du mois, puis soustraire la valeur du "
         "Monthly."),
        ("4. Choisir l'export de référence",
         "Garde le Daily, écarte le Monthly, n'utilise le Hourly que pour les heures.",
         "Travailler uniquement à partir de l'onglet Données."),
        ("5. Recalculer année, mois, jour",
         "Ignore les colonnes Year, Month, Hour et Weekday Name des exports et les recalcule à partir de la date.",
         "ANNEE(date), MOIS(date), et TEXTE(date;\"jjjj\") pour le nom du jour."),
        ("6. Repérer les mois complets",
         "Un mois est complet si tous ses jours sont présents. Les autres sont exclus des statistiques mensuelles.",
         "NB.SI.ENS sur les dates du mois, comparé à JOUR(FIN.MOIS(date;0))."),
        ("7. Statistiques par groupe",
         "Moyenne, médiane, écart-type, maximum par jour ; coefficient de variation = écart-type / moyenne ; part "
         "des jours sans vente.",
         "MOYENNE, MEDIANE, ECARTYPE.STANDARD, MAX sur la colonne du groupe ; NB.SI(colonne;0)/NBVAL(colonne)."),
        ("8. Saisonnalité",
         "Moyenne des ventes de chaque mois de l'année sur les mois complets. Indice = cette moyenne divisée par "
         "la moyenne de tous les mois complets.",
         "Sur le tableau des mois complets de l'onglet Statistiques : MOYENNE.SI sur le numéro du mois, puis "
         "division par MOYENNE de la colonne."),
        ("9. Calendrier et conseils",
         f"Un mois est dit fort si son indice atteint {SEUIL_FORT}, creux s'il descend à {SEUIL_CREUX}. La "
         "régularité compte, sur les années entières, combien de fois le mois est passé sous la moyenne de son "
         "année. L'évolution sur un an compare les 12 derniers mois complets aux 12 précédents.",
         "Mise en forme conditionnelle sur les indices ; pour l'évolution, SOMME des 12 derniers mois complets "
         "divisée par SOMME des 12 précédents, moins 1."),
        ("10. Jour de semaine et heure",
         "Moyenne des ventes par jour de la semaine (Daily) et par heure (Hourly).",
         "MOYENNE.SI(colonne Jour;\"Lundi\";colonne du groupe). Les heures demandent d'ouvrir le fichier Hourly."),
        ("11. Évolution par année",
         "Moyenne par jour pour chaque année, et non un total, car la première et la dernière année sont incomplètes.",
         "MOYENNE.SI(colonne Année;2018;colonne du groupe)."),
        ("12. Onglet Ventes",
         "Écrit des formules SOMME.SI.ENS qui additionnent l'onglet Données entre la date de début et la date de fin.",
         "Les formules sont visibles dans l'onglet Ventes : cliquer sur une cellule de résultat pour la lire."),
        ("13. Mise à jour mensuelle",
         "Remplacer les 4 CSV dans le dossier data, lancer python src/build_workbook.py : les deux classeurs sont "
         "recréés et tous les contrôles refaits.",
         "Coller les nouvelles lignes à la suite dans l'onglet Données, puis étendre les plages des formules de "
         "l'onglet Ventes. Les autres onglets ne se recalculent pas seuls."),
    ]
    return pd.DataFrame(lignes, columns=["Étape", "Ce que fait le programme", "Comment le refaire à la main dans Excel"])


def faq(exports, mois):
    d = exports["Daily"]
    incomplets = ", ".join(mois.loc[mois["Complet"] == "non", "Mois"])
    nb_ecarts = int((mois["Monthly en écart"] == "oui").sum())
    lignes = [
        ("En quelle unité sont les quantités ?",
         "On ne le sait pas : le logiciel de caisse ne le documente pas. Les quantités ont des décimales, ce ne "
         "sont donc pas des boîtes entières."),
        ("Puis-je additionner ou comparer deux groupes ?",
         "Non. Faute d'unité connue, rien ne dit qu'une unité d'un groupe vaut une unité d'un autre. Comparez un "
         "groupe à lui-même dans le temps."),
        ("Que contient exactement chaque groupe ?",
         "L'export regroupe les ventes par code ATC, la classification internationale des médicaments tenue par "
         "l'OMS. Le nom officiel de chaque groupe et des exemples de substances sont dans l'onglet Médicaments."),
        ("Puis-je connaître le chiffre d'affaires ou la marge ?",
         "Non, l'export ne contient aucun prix."),
        ("Pourquoi les mois de ce classeur ne correspondent-ils pas à mon export mensuel ?",
         f"L'export mensuel ne concorde pas avec les ventes jour par jour sur {nb_ecarts} mois (liste dans "
         "l'onglet État des données). Les exports horaire, journalier et hebdomadaire concordent entre eux : les "
         "mois sont donc recalculés à partir des jours."),
        ("Pourquoi certains mois manquent-ils dans les statistiques mensuelles ?",
         f"Les mois incomplets sont exclus pour ne pas fausser les moyennes : {incomplets}. Leurs jours restent "
         "disponibles dans l'onglet Ventes."),
        ("Comment voir les ventes d'un seul jour ?",
         "Dans l'onglet Ventes, mettez la même date en début et en fin."),
        ("Pourquoi la colonne Évolution est-elle vide ?",
         "Soit le groupe est sur NON, soit il n'y a aucune vente à la même période un an plus tôt (par exemple "
         "pour une période de 2014)."),
        ("Quelles dates puis-je choisir ?",
         f"Du {d['datum'].min():%d/%m/%Y} au {d['datum'].max():%d/%m/%Y}. Une date en dehors est refusée."),
        ("Que veulent dire les couleurs du calendrier ?",
         f"Orange : mois fort, au moins {SEUIL_FORT - 1:.0%} au-dessus d'un mois moyen du groupe. Bleu : mois "
         f"creux, au moins {1 - SEUIL_CREUX:.0%} en dessous."),
        ("Puis-je modifier l'onglet Données ?",
         "Il vaut mieux ne pas y toucher : l'onglet Ventes lit ses chiffres dedans. Ce sont les ventes telles "
         "qu'exportées, sans aucune modification."),
        ("Comment mettre à jour le classeur le mois prochain ?",
         "Envoyez les quatre nouveaux exports : le programme recrée le classeur et refait tous les contrôles. "
         "La marche à suivre manuelle est décrite dans l'onglet Méthode."),
    ]
    return pd.DataFrame(lignes, columns=["Question", "Réponse"])


def statistiques(exports):
    d = exports["Daily"].set_index("datum")[G]
    h = exports["Hourly"].set_index("datum")[G]
    complets = mois_complets(d)
    mensuel = d.resample("ME").sum().loc[complets]

    par_groupe = pd.DataFrame({
        "Nom": pd.Series(GROUPES),
        "Total": d.sum().round(1),
        "Moyenne par jour": d.mean().round(2),
        "Médiane par jour": d.median().round(2),
        "Écart-type par jour": d.std().round(2),
        "Coefficient de variation": (d.std() / d.mean()).round(2),
        "Maximum sur un jour": d.max().round(1),
        "Part des jours sans vente": (d == 0).mean().round(3),
        "Moyenne par mois complet": mensuel.mean().round(1),
        "Mois le plus bas": mensuel.min().round(1),
        "Mois le plus haut": mensuel.max().round(1),
    }).rename_axis("Groupe").reset_index()

    saison = mensuel.groupby(mensuel.index.month).mean()
    indice = (saison / mensuel.mean()).round(2)
    saison.index = indice.index = [MOIS[i - 1] for i in saison.index]

    semaine = d.groupby(d.index.dayofweek).mean()
    semaine.index = JOURS

    heures = h.groupby(h.index.hour).mean()
    heures = heures[heures.sum(axis=1) > 0]
    heures.index = [f"{i}h" for i in heures.index]

    annees = d.groupby(d.index.year).mean()
    annees.insert(0, "Jours de données", d.groupby(d.index.year).size())

    serie = mensuel.copy()
    serie.index = serie.index.strftime("%Y-%m")

    return {
        "groupes": par_groupe,
        "saison": saison.round(1).rename_axis("Mois").reset_index(),
        "indice": indice.rename_axis("Mois").reset_index(),
        "semaine": semaine.round(2).rename_axis("Jour").reset_index(),
        "heures": heures.round(3).rename_axis("Heure").reset_index(),
        "annees": annees.round(2).rename_axis("Année").reset_index(),
        "mensuel": serie.round(2).rename_axis("Mois complet").reset_index(),
    }


def saisons(exports, stats):
    """Calendrier saisonnier : indice par mois, régularité d'une année à l'autre, conseil par groupe."""
    d = exports["Daily"].set_index("datum")[G]
    mensuel = d.resample("ME").sum().loc[mois_complets(d)]
    indice = stats["indice"].set_index("Mois").T  # groupes x mois

    # Années entières : combien de fois chaque mois est passé sous la moyenne de son année
    entieres = [a for a, x in mensuel.groupby(mensuel.index.year) if len(x) == 12]
    m = mensuel[mensuel.index.year.isin(entieres)]
    sous = (m < m.groupby(m.index.year).transform("mean")).groupby(m.index.month).sum().T
    sous.columns = MOIS

    dernier, avant = mensuel.iloc[-12:], mensuel.iloc[-24:-12]
    tendance = dernier.sum() / avant.sum() - 1
    amplitude = indice.max(axis=1) - indice.min(axis=1)
    sans_vente = (d == 0).mean()
    semaine = stats["semaine"].set_index("Jour")

    lignes = []
    for g in G:
        forts = [mo for mo in MOIS if indice.loc[g, mo] >= SEUIL_FORT]
        creux = [mo for mo in MOIS if indice.loc[g, mo] <= SEUIL_CREUX]
        force = "Forte" if amplitude[g] >= 0.6 else "Modérée" if amplitude[g] >= 0.3 else "Faible"
        rare = sans_vente[g] > 0.5
        if rare:
            conseil = (f"Ventes rares (aucune vente {sans_vente[g]:.0%} des jours) : stock minimal, "
                       "réassort à la demande.")
        elif force != "Faible" and forts:
            conseil = f"Renforcer le stock avant {', '.join(forts).lower()}"
            conseil += f" ; l'alléger en {', '.join(creux).lower()}." if creux else "."
        else:
            conseil = "Demande régulière toute l'année : stock constant."
        if tendance[g] > 0.10:
            conseil += f" Demande en hausse sur un an ({tendance[g]:+.0%}) : prévoir plus que l'an dernier."
        elif tendance[g] < -0.10:
            conseil += f" Demande en baisse sur un an ({tendance[g]:+.0%}) : prévoir moins que l'an dernier."
        else:
            conseil += " Demande stable sur un an."
        lignes.append({
            "Groupe": GROUPES[g],
            "Code": g,
            "Conseil de commande": conseil,
            "Mois forts": ", ".join(forts) or "aucun",
            "Mois creux": ", ".join(creux) or "aucun",
            "Évolution sur un an": f"{tendance[g]:+.0%}",
            "Saisonnalité": "Non significative (ventes rares)" if rare else force,
            "Jour le plus fort": semaine[g].idxmax(),
            "Jour le plus faible": semaine[g].idxmin(),
        })
    periode = (f"{dernier.index[0]:%m/%Y} à {dernier.index[-1]:%m/%Y} comparé à "
               f"{avant.index[0]:%m/%Y} à {avant.index[-1]:%m/%Y}")
    return {"indice": indice, "sous": sous, "annees": entieres, "conseils": pd.DataFrame(lignes),
            "tendance": tendance, "amplitude": amplitude, "rares": [g for g in G if sans_vente[g] > 0.5],
            "sans_vente": sans_vente, "moy_dernier": dernier.mean(), "moy_avant": avant.mean(),
            "periode": periode, "dernier_mois": mensuel.index[-1]}


def donnees_jour(exports):
    d = exports["Daily"][["datum"] + G].rename(columns={"datum": "Date"})
    d["Année"] = d["Date"].dt.year
    d["Mois"] = d["Date"].dt.month
    d["Jour"] = [JOURS[i] for i in d["Date"].dt.dayofweek]
    return d


# Mise en forme : fonds pastel, texte foncé, bordures fines. Chaque page prend la couleur de son onglet.
DEFAUT = ("", "E7E6E6", "404040")
BORD = Border(*(Side(style="thin", color="D9D9D9"),) * 4)
SAISIE = PatternFill("solid", fgColor="FFF2CC")
FORT = PatternFill("solid", fgColor="F4B183")
CREUX = PatternFill("solid", fgColor="BDD7EE")
OUI = PatternFill("solid", fgColor="C6EFCE")
NON = PatternFill("solid", fgColor="FFC7CE")
ATTENTION = PatternFill("solid", fgColor="FFEB9C")
# Mots qui portent toujours la même couleur, partout dans le classeur
COULEURS = {"OUI": OUI, "OK": OUI, "NON": NON, "À CONFIRMER": ATTENTION, "À noter": ATTENTION,
            "Mois forts": FORT, "Mois creux": CREUX}
MILIEU = Alignment(vertical="center", wrap_text=True)
CENTRE = Alignment(horizontal="center", vertical="center", wrap_text=True)
SANS_RENVOI = Alignment(vertical="center")


def clair(onglet):
    return PatternFill("solid", fgColor=ONGLETS.get(onglet, DEFAUT)[1])


def fonce(onglet):
    return ONGLETS.get(onglet, DEFAUT)[2]


def ecrire(ws, r, c1, valeur, c2=None, font=None, fill=None, alignment=MILIEU, bord=True, fmt=None):
    """Écrit une valeur, sur une cellule ou une plage fusionnée, et ajuste la hauteur de ligne au texte."""
    c2 = c2 or c1
    for c in range(c1, c2 + 1):
        x = ws.cell(row=r, column=c)
        if bord:
            x.border = BORD
        if fill:
            x.fill = fill
    cell = ws.cell(row=r, column=c1, value=valeur)
    cell.alignment = alignment
    if font:
        cell.font = font
    if fmt:
        cell.number_format = fmt
    if c2 > c1:
        ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
    if isinstance(valeur, str) and alignment.wrap_text:
        largeur = sum(ws.column_dimensions[get_column_letter(c)].width for c in range(c1, c2 + 1))
        taille = font.size if font and font.size else 11
        lignes = sum(max(1, ceil(len(p) * taille / 11 / (largeur * 0.92))) for p in valeur.split("\n"))
        ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, lignes * taille * 1.3 + 6)
    return cell


def bouton(ws, r, c1, onglet, texte=None, c2=None, retour=False):
    """Bouton cliquable à la couleur de l'onglet qu'il ouvre."""
    libelle = f"◀ {texte or onglet}" if retour else f"{texte or onglet} ▶"
    cell = ecrire(ws, r, c1, libelle, c2, font=Font(bold=True, color=fonce(onglet)), fill=clair(onglet),
                  alignment=CENTRE)
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location=f"'{onglet}'!A1")
    cadre = Border(*(Side(style="thin", color=fonce(onglet)),) * 4)
    for c in range(c1, (c2 or c1) + 1):
        ws.cell(row=r, column=c).border = cadre
    ws.row_dimensions[r].height = max(ws.row_dimensions[r].height or 15, 28)
    return cell


def banniere(ws, titre, sous_titre, nb_col):
    """Titre de la page sur un bandeau à la couleur de l'onglet, puis bouton de retour à l'accueil."""
    ws.sheet_view.showGridLines = False
    ecrire(ws, 1, 1, titre, nb_col, font=Font(bold=True, size=18, color=fonce(ws.title)), fill=clair(ws.title),
           bord=False)
    ws.row_dimensions[1].height = 38
    if sous_titre:
        ecrire(ws, 2, 1, sous_titre, nb_col, font=Font(italic=True, color="404040"), bord=False)
    if ws.title != "Synthèse":
        bouton(ws, 3, 1, "Synthèse", "Accueil", retour=True)


def section(ws, r, texte, col=1):
    ecrire(ws, r, col, texte.upper(), font=Font(bold=True, size=12, color=fonce(ws.title)), bord=False,
           alignment=SANS_RENVOI)
    ws.row_dimensions[r].height = 24


def entetes(ws, r, textes, debut=1, fusions=None):
    """Ligne de titres de colonnes. fusions : {position dans textes: nombre de colonnes occupées}."""
    c = debut
    for k, texte in enumerate(textes):
        n = (fusions or {}).get(k, 1)
        ecrire(ws, r, c, texte, c + n - 1, font=Font(bold=True, color=fonce(ws.title)),
               fill=COULEURS.get(texte, clair(ws.title)), alignment=CENTRE)
        c += n


def legende(ws, r, col, lignes):
    """Légende : une case colorée portant son nom, et à droite ce qu'elle signifie."""
    section(ws, r, "Légende", col)
    for k, (fill, nom, sens) in enumerate(lignes, start=1):
        ecrire(ws, r + k, col, nom, fill=fill, font=Font(bold=True), alignment=CENTRE)
        ecrire(ws, r + k, col + 1, sens, bord=False, alignment=SANS_RENVOI)
    return r + len(lignes) + 1


def ecrire_tableaux(ws, titre, sous_titre, blocs, largeur_max=60):
    """Écrit un ou plusieurs tableaux l'un sous l'autre, chacun avec son titre."""
    largeurs = {}
    for _, df in blocs:
        for j, col in enumerate(df.columns, start=1):
            contenu = [len(str(col))] + [len(str(v)) for v in df[col]]
            largeurs[j] = max(largeurs.get(j, 0), max(contenu))
    for j, largeur in largeurs.items():
        ws.column_dimensions[get_column_letter(j)].width = min(max(largeur + 2, 12), largeur_max)
    banniere(ws, titre, sous_titre, max(len(largeurs), 4))
    r = 5
    for sous, df in blocs:
        if sous:
            section(ws, r, sous)
            r += 1
        entetes(ws, r, list(df.columns))
        if len(blocs) == 1:
            ws.freeze_panes = ws.cell(row=r + 1, column=1)
        for ligne in df.itertuples(index=False):
            r += 1
            for j, v in enumerate(ligne, start=1):
                ecrire(ws, r, j, v, fill=COULEURS.get(v) if isinstance(v, str) else None)
        r += 3


def ecrire_synthese(ws, exports, controles, mois, stats, s):
    d = exports["Daily"].set_index("datum")[G]
    complets = mois_complets(d)
    a_prevoir = complets.max() + pd.offsets.MonthEnd(1)
    prochain = MOIS[a_prevoir.month - 1]
    res = controles.set_index("Contrôle")["Résultat"]
    anomalies = controles[controles["Contrôle"].str.contains("valeurs manquantes") & (controles["Statut"] != "OK")]
    incomplets = ", ".join(mois.loc[mois["Complet"] == "non", "Mois"]) or "aucun"
    nb_ecarts = int((mois["Monthly en écart"] == "oui").sum())
    ind, sous, n_ans = s["indice"], s["sous"], len(s["annees"])
    tendance, amplitude, rares = s["tendance"], s["amplitude"], s["rares"]
    semaine = stats["semaine"].set_index("Jour")
    moi = "Synthèse"

    def noms(codes):
        return ", ".join(GROUPES[g] for g in codes) or "aucun"

    def fr(x, nb=2):
        return f"{x:.{nb}f}".replace(".", ",")

    for lettre, largeur in zip("ABCDEFGH", (5, 62, 15, 32, 14, 32, 25, 26)):
        ws.column_dimensions[lettre].width = largeur
    banniere(ws, "Ventes de la pharmacie : synthèse pour mieux commander",
             "Question du client : anticiper la demande. Quels médicaments, quels jours, quelles périodes de l'année ?", 8)

    autres = [o for o in ONGLETS if o != moi]
    ecrire(ws, 4, 2, "Aller à :", font=Font(bold=True, color=fonce(moi)), bord=False,
           alignment=Alignment(horizontal="right", vertical="center"))
    for k, onglet in enumerate(autres):
        bouton(ws, 4 + k // 6, 3 + k % 6, onglet)
    r = 7

    section(ws, r, f"À préparer pour {prochain.lower()}")
    for libelle, codes, fill in (
            ("À renforcer", [g for g in G if ind.loc[g, prochain] >= SEUIL_FORT and g not in rares], FORT),
            ("À alléger", [g for g in G if ind.loc[g, prochain] <= SEUIL_CREUX and g not in rares], CREUX)):
        r += 1
        ecrire(ws, r, 1, libelle, 2, font=Font(bold=True), fill=fill)
        ecrire(ws, r, 3, noms(codes), 8, font=Font(bold=True, size=12))
    r += 2

    section(ws, r, "D'où viennent les chiffres")
    r += 1
    ecrire(ws, r, 1, f"① Une seule source : l'export journalier. Les exports horaire et hebdomadaire disent "
                     f"exactement la même chose.\n② Aucune vente corrigée : ni case vide, ni doublon, ni valeur "
                     f"négative sur {len(d)} jours.", 2, fill=clair(moi), bord=False)
    ecrire(ws, r, 3, f"{nb_ecarts} / {len(mois)}", font=Font(bold=True, size=16, color="C00000"), alignment=CENTRE,
           bord=False)
    ecrire(ws, r, 4, "mois où l'export mensuel annonce plus de ventes que les jours. Il est écarté : les mois sont "
                     "recalculés à partir des jours.", 7, bord=False)
    bouton(ws, r, 8, "État des données", "Voir le détail")
    r += 1
    ecrire(ws, r, 1, f"Portée : {len(complets)} mois complets, de {complets.min():%m/%Y} à {complets.max():%m/%Y} "
                     f"(mois incomplets exclus des moyennes : {incomplets}). L'unité des quantités n'est pas "
                     "documentée par le logiciel de caisse : chaque groupe est comparé à lui-même dans le temps, "
                     "jamais aux autres.", 8, bord=False)
    r += 2

    # Conclusions : (catégorie, phrase, chiffre, ce que c'est, comparé à, ce que c'est, fiabilité, onglet de preuve)
    conclusions = []
    for g in sorted((g for g in G if amplitude[g] >= 0.6 and g not in rares), key=lambda g: -amplitude[g]):
        haut, bas = ind.loc[g].idxmax(), ind.loc[g].idxmin()
        fois = int(sous.loc[g, bas])
        conclusions.append((
            "fiable" if fois >= n_ans - 1 else "attention",
            f"{GROUPES[g]} : forte saison. {fr(ind.loc[g, haut] / ind.loc[g, bas], 1)} fois plus de ventes en "
            f"{haut.lower()} qu'en {bas.lower()}.",
            fr(ind.loc[g, haut]), f"Indice de {haut.lower()} (1 = mois moyen du groupe)",
            fr(ind.loc[g, bas]), f"Indice de {bas.lower()}",
            f"Oui : creux en {bas.lower()} {fois} années sur {n_ans}", "Calendrier"))
    for g in (g for g in G if tendance[g] > 0.10 and g not in rares):
        conclusions.append((
            "attention", f"{GROUPES[g]} : la demande monte, {tendance[g]:+.0%} en un an.",
            fr(s["moy_dernier"][g], 0), "Ventes moyennes par mois, 12 derniers mois complets",
            fr(s["moy_avant"][g], 0), "Ventes moyennes par mois, 12 mois précédents",
            "À suivre : une seule année de recul", "Conseils"))
    stables = [g for g in G if amplitude[g] < 0.3]
    if stables:
        top = amplitude.idxmax()
        conclusions.append((
            "fiable", f"{noms(stables)} : demande régulière toute l'année, stock constant.",
            fr(amplitude[stables].max()), "Plus grand écart d'indice entre mois fort et mois faible (ces groupes)",
            fr(amplitude[top]), f"Même écart pour {GROUPES[top].lower()}", "Oui", "Calendrier"))
    for g in rares:
        mini = s["sans_vente"].idxmin()
        conclusions.append((
            "info", f"{GROUPES[g]} : ventes rares, aucune saison fiable. Stock minimal, réassort à la demande.",
            f"{s['sans_vente'][g]:.0%}", "Jours sans aucune vente",
            f"{s['sans_vente'][mini]:.0%}", f"Jours sans vente, {GROUPES[mini].lower()}", "Oui", "Statistiques"))
    autres_jours = semaine.drop("Dimanche").mean()
    g = (semaine.loc["Dimanche"] / autres_jours).drop(rares).idxmin()
    if semaine.loc["Dimanche", g] / autres_jours[g] < 0.8:
        conclusions.append((
            "fiable", f"{GROUPES[g]} : nettement moins de ventes le dimanche.",
            fr(semaine.loc["Dimanche", g], 1), "Ventes moyennes un dimanche",
            fr(autres_jours[g], 1), "Ventes moyennes un autre jour",
            f"Oui : moyenne sur {int((d.index.dayofweek == 6).sum())} dimanches", "Statistiques"))
    conclusions.append((
        "attention", "L'export mensuel ne concorde pas avec les ventes jour par jour : il n'est pas utilisé.",
        str(nb_ecarts), "Mois en écart", str(len(mois)), "Mois dans l'export",
        "À signaler à l'éditeur du logiciel de caisse", "État des données"))
    conclusions.append((
        "info", f"Peut-on prévoir les ventes du mois suivant ({a_prevoir:%m/%Y}) ? Le test n'est pas encore réalisé.",
        "?", "Erreur de la prévision testée", "?", "Erreur de la prévision la plus simple", "À CONFIRMER", None))

    teintes = {"fiable": PatternFill("solid", fgColor="D5F0E1"), "attention": PatternFill("solid", fgColor="FAD4E0"),
               "info": PatternFill("solid", fgColor="DCE8FA")}
    section(ws, r, "Les conclusions et les chiffres qui les prouvent")
    r += 1
    entetes(ws, r, ["N°", "Conclusion", "Chiffre", "Ce que c'est", "Comparé à", "Ce que c'est",
                    "Peut-on s'y fier ?", "Voir la preuve"])
    for k, (categorie, phrase, chiffre, sens, compare, sens2, fiable, preuve) in enumerate(conclusions, start=1):
        r += 1
        ecrire(ws, r, 1, k, fill=teintes[categorie], alignment=CENTRE)
        ecrire(ws, r, 2, phrase, fill=teintes[categorie], font=Font(bold=True))
        ecrire(ws, r, 3, chiffre, font=Font(bold=True, size=14, color=fonce(moi)), alignment=CENTRE)
        ecrire(ws, r, 4, sens)
        ecrire(ws, r, 5, compare, font=Font(bold=True, size=14, color="7F7F7F"), alignment=CENTRE)
        ecrire(ws, r, 6, sens2)
        ecrire(ws, r, 7, fiable, fill=COULEURS.get(fiable, teintes[categorie]))
        if preuve:
            bouton(ws, r, 8, preuve)
        else:
            ecrire(ws, r, 8, "")
        ws.row_dimensions[r].height = max(ws.row_dimensions[r].height, 40)
    r += 1
    r = legende(ws, r + 1, 2, [
        (teintes["fiable"], "Vert", "Conclusion solide : elle se répète d'une année à l'autre."),
        (teintes["attention"], "Rose", "Point à surveiller ou à faire confirmer."),
        (teintes["info"], "Bleu", "Information, ou réponse pas encore établie."),
    ]) + 1

    section(ws, r, "Traitement des données : ce qui a été vérifié")
    r += 1
    entetes(ws, r, ["", "Point vérifié", "Résultat", "Constat", "Décision"], fusions={3: 3, 4: 2})
    propre = "OK" if anomalies.empty else "À noter"
    for point, statut, constat, decision in (
        ("Qualité de chaque fichier", propre,
         "Aucune case vide, aucune date en double, aucune quantité négative" if anomalies.empty
         else "Anomalies détectées : voir l'onglet État des données",
         "Aucune correction nécessaire" if anomalies.empty else "À examiner avant d'utiliser les chiffres"),
        ("Jours et heures manquants", "OK",
         f"{res['Daily : jours absents entre le premier et le dernier jour']} jour absent, "
         f"{res['Hourly : heures absentes entre la première et la dernière heure']} heure absente",
         "Historique continu, utilisable tel quel"),
        ("Horaire comparé au journalier", "OK", res["Hourly regroupé par jour = Daily"],
         "L'horaire sert uniquement aux profils par heure"),
        ("Journalier comparé à l'hebdomadaire", "OK", res["Daily regroupé par semaine = Weekly"],
         "L'hebdomadaire n'apporte rien de plus : non utilisé"),
        ("Journalier comparé au mensuel", "À noter", res["Daily regroupé par mois = Monthly"],
         "Mensuel écarté ; les mois sont recalculés à partir des jours"),
        ("Mois incomplets", "À noter", incomplets, "Exclus des moyennes mensuelles, conservés dans l'onglet Ventes"),
        ("Colonnes Year, Month, Hour, Weekday Name", "À noter",
         "Redondantes avec la date ; « Hour » du journalier est une somme sans signification",
         "Ignorées et recalculées à partir de la date"),
        ("Unité des quantités", "À noter", "Non documentée ; les quantités ont des décimales",
         "Les groupes ne sont jamais additionnés ni comparés entre eux"),
        ("Valeurs de vente", "OK", "Aucune erreur de saisie démontrable", "Aucune vente modifiée ni supprimée"),
    ):
        r += 1
        ecrire(ws, r, 1, "")
        ecrire(ws, r, 2, point, font=Font(bold=True))
        ecrire(ws, r, 3, statut, fill=COULEURS[statut], alignment=CENTRE, font=Font(bold=True))
        ecrire(ws, r, 4, constat, 6)
        ecrire(ws, r, 7, decision, 8)
    r += 1
    ecrire(ws, r, 1, ("Conclusion du traitement : les données sont utilisables telles quelles. " if anomalies.empty
                      else "Conclusion du traitement : des anomalies sont à examiner avant d'utiliser les chiffres. ")
           + "Une seule source est retenue (le journalier), un export est écarté (le mensuel), et aucune vente n'a "
             "été modifiée.", 8, font=Font(bold=True), fill=clair(moi), bord=False)
    r += 2

    section(ws, r, "Ce que les chiffres peuvent dire, et ce qu'ils ne peuvent pas dire")
    r += 1
    entetes(ws, r, ["", "Question", "Réponse", "Pourquoi"], fusions={3: 5})
    for question, reponse, pourquoi in (
        ("Quel est le niveau habituel des ventes d'un groupe ?", "OUI", "Ventes connues jour par jour sur près de six ans."),
        ("Quelles périodes de l'année sont fortes ou faibles ?", "OUI", "Voir l'onglet Calendrier."),
        ("Quels jours et quelles heures sont les plus chargés ?", "OUI", "Voir l'onglet Statistiques."),
        ("Un groupe monte-t-il ou baisse-t-il d'une année à l'autre ?", "OUI", "Voir l'onglet Conseils."),
        ("Quel est le chiffre d'affaires, quelle est la marge ?", "NON", "L'export ne contient aucun prix."),
        ("Quel groupe se vend plus qu'un autre ?", "NON", "L'unité des quantités n'est pas documentée."),
        ("Quel médicament précis se vend ?", "NON", "Les ventes sont regroupées par famille (code ATC)."),
        ("Qui achète, combien de patients ?", "NON", "L'export ne contient aucun identifiant de client."),
        ("Pourquoi les ventes varient-elles ?", "NON", "Il faut les croiser avec des données extérieures."),
        (f"Peut-on prévoir le mois suivant ({a_prevoir:%m/%Y}) ?", "À CONFIRMER",
         "Le test de prévision n'est pas encore réalisé."),
    ):
        r += 1
        ecrire(ws, r, 1, "")
        ecrire(ws, r, 2, question, font=Font(bold=True))
        ecrire(ws, r, 3, reponse, fill=COULEURS[reponse], alignment=CENTRE, font=Font(bold=True))
        ecrire(ws, r, 4, pourquoi, 8)
    r += 3

    section(ws, r, "Sommaire")
    for onglet in autres:
        r += 1
        bouton(ws, r, 1, onglet, c2=2)
        ecrire(ws, r, 3, ONGLETS[onglet][0], 8)
    legende(ws, r + 2, 2, [
        (clair("Ventes"), "Case ▶", "Bouton cliquable. Sa couleur est celle de l'onglet qu'il ouvre."),
        (OUI, "OUI / OK", "Les chiffres permettent de répondre ; le contrôle est bon."),
        (NON, "NON", "Les chiffres ne permettent pas de répondre."),
        (ATTENTION, "À noter", "Point à connaître, ou réponse pas encore établie."),
    ])


def ecrire_ventes(ws, jour):
    """Onglet de sélection : formules SOMME.SI.ENS sur l'onglet Données, sans macro."""
    n = len(jour) + 1
    dates = f"'Données'!$A$2:$A${n}"
    fin = jour["Date"].max()
    debut = fin.replace(day=1)
    for lettre, largeur in zip("ABCDEFG", (38, 16, 20, 18, 20, 14, 60)):
        ws.column_dimensions[lettre].width = largeur
    banniere(ws, "Ventes sur un jour ou une période",
             "Modifiez les cases jaunes. Pour un seul jour, mettez la même date au début et à la fin. "
             f"Données disponibles du {jour['Date'].min():%d/%m/%Y} au {fin:%d/%m/%Y}.", 6)

    for r, (libelle, valeur) in enumerate([("Date de début", debut), ("Date de fin", fin)], start=5):
        ecrire(ws, r, 1, libelle, font=Font(bold=True))
        ecrire(ws, r, 2, valeur.to_pydatetime(), fill=SAISIE, alignment=CENTRE, fmt="DD/MM/YYYY")
    ecrire(ws, 7, 1, "Nombre de jours", font=Font(bold=True))
    ecrire(ws, 7, 2, '=IF($B$6<$B$5,"Fin avant le début",$B$6-$B$5+1)', alignment=CENTRE)

    entetes(ws, 9, ["Groupe", "Afficher ?", "Ventes sur la période", "Moyenne par jour",
                    "Même période un an plus tôt", "Évolution"])
    for i, g in enumerate(G):
        r = 10 + i
        col = f"'Données'!${get_column_letter(2 + i)}$2:${get_column_letter(2 + i)}${n}"
        ecrire(ws, r, 1, GROUPES[g], font=Font(bold=True))
        ecrire(ws, r, 2, "OUI", fill=SAISIE, alignment=CENTRE)
        ecrire(ws, r, 3, f'=IF($B{r}="OUI",SUMIFS({col},{dates},">="&$B$5,{dates},"<="&$B$6),"")',
               alignment=CENTRE, fmt="0.0", font=Font(bold=True, size=12, color=fonce("Ventes")))
        ecrire(ws, r, 4, f'=IF(AND($B{r}="OUI",ISNUMBER($B$7)),C{r}/$B$7,"")', alignment=CENTRE, fmt="0.00")
        ecrire(ws, r, 5, f'=IF($B{r}="OUI",SUMIFS({col},{dates},">="&EDATE($B$5,-12),'
                         f'{dates},"<="&EDATE($B$6,-12)),"")', alignment=CENTRE, fmt="0.0")
        ecrire(ws, r, 6, f'=IF(AND($B{r}="OUI",N(E{r})>0),C{r}/E{r}-1,"")', alignment=CENTRE, fmt="+0%;-0%;0%")
        ws.row_dimensions[r].height = 22

    choix = DataValidation(type="list", formula1='"OUI,NON"', allow_blank=False)
    choix.add(f"B10:B{9 + len(G)}")
    bornes = DataValidation(type="date", operator="between",
                            formula1=f"MIN({dates})", formula2=f"MAX({dates})",
                            showErrorMessage=True, errorTitle="Date hors des données",
                            error="Choisissez une date comprise dans la période disponible.")
    bornes.add("B5:B6")
    ws.add_data_validation(choix)
    ws.add_data_validation(bornes)

    bas = 11 + len(G)
    ecrire(ws, bas, 1, "Comparez un groupe à lui-même dans le temps, pas les groupes entre eux : l'unité des "
                       "quantités n'est pas connue.", 6, font=Font(italic=True, color="555555"), bord=False)
    bouton(ws, bas + 2, 1, "Calendrier", "Voir le calendrier des saisons")
    bouton(ws, bas + 2, 3, "Conseils", "Lire les conseils de commande", c2=4)
    legende(ws, bas + 4, 1, [
        (SAISIE, "Case jaune", "À vous de la modifier : les deux dates, et OUI ou NON pour chaque groupe."),
        (None, "Case blanche", "Résultat calculé automatiquement : ne pas modifier."),
        (clair("Calendrier"), "Case ▶", "Bouton cliquable. Sa couleur est celle de l'onglet qu'il ouvre."),
    ])


def ecrire_calendrier(ws, s):
    """Calendrier saisonnier coloré : mois forts et mois creux de chaque groupe."""
    indice, sous, annees = s["indice"], s["sous"], s["annees"]
    ws.column_dimensions["A"].width = 38
    for j in range(2, 14):
        ws.column_dimensions[get_column_letter(j)].width = 11
    banniere(ws, "Calendrier des saisons", "Pour chaque groupe, les mois où il se vend plus ou moins que d'habitude.", 13)

    def couleur(g, mo):
        v = indice.loc[g, mo]
        return FORT if v >= SEUIL_FORT else CREUX if v <= SEUIL_CREUX else None

    def grille(r, valeur, fmt=None):
        entetes(ws, r, ["Groupe"] + MOIS)
        for i, g in enumerate(G, start=1):
            ecrire(ws, r + i, 1, GROUPES[g], font=Font(bold=True))
            for j, mo in enumerate(MOIS, start=2):
                ecrire(ws, r + i, j, valeur(g, mo), fill=couleur(g, mo), alignment=CENTRE, fmt=fmt)
            ws.row_dimensions[r + i].height = 22

    r = legende(ws, 5, 1, [
        (FORT, "Mois fort", f"Ventes au moins {SEUIL_FORT - 1:.0%} au-dessus d'un mois moyen du groupe : renforcer le stock."),
        (CREUX, "Mois creux", f"Ventes au moins {1 - SEUIL_CREUX:.0%} en dessous d'un mois moyen : alléger le stock."),
        (None, "Mois normal", "Ventes proches de la moyenne. Le chiffre 1,00 correspond à un mois moyen du groupe."),
    ]) + 1
    section(ws, r, "Le calendrier")
    grille(r + 1, lambda g, mo: float(indice.loc[g, mo]), "0.00")

    r += 4 + len(G)
    section(ws, r, "À préparer pour les trois prochains mois")
    entetes(ws, r + 1, ["Mois", "À renforcer", "À alléger"], fusions={1: 6, 2: 6})
    for k in range(1, 4):
        mo = MOIS[(s["dernier_mois"] + pd.offsets.MonthEnd(k)).month - 1]
        ecrire(ws, r + 1 + k, 1, mo, font=Font(bold=True))
        ecrire(ws, r + 1 + k, 2, ", ".join(
            GROUPES[g] for g in G if indice.loc[g, mo] >= SEUIL_FORT and g not in s["rares"]) or "rien de particulier", 7)
        ecrire(ws, r + 1 + k, 8, ", ".join(
            GROUPES[g] for g in G if indice.loc[g, mo] <= SEUIL_CREUX and g not in s["rares"]) or "rien de particulier", 13)

    r += 7
    section(ws, r, f"Peut-on s'y fier ? Nombre d'années, sur {len(annees)} ({annees[0]} à {annees[-1]}), où le mois "
                   "a été sous la moyenne de son année")
    grille(r + 1, lambda g, mo: f"{int(sous.loc[g, mo])} sur {len(annees)}")
    bouton(ws, r + len(G) + 3, 1, "Ventes", "Voir les ventes d'une période")


def blocs_statistiques(stats):
    return [
        ("Par groupe (ventes par jour, puis par mois complet)", stats["groupes"]),
        ("Indice saisonnier (1 = mois moyen du groupe)", stats["indice"]),
        ("Ventes moyennes de chaque mois de l'année", stats["saison"]),
        ("Ventes moyennes par jour de la semaine", stats["semaine"]),
        ("Ventes moyennes par heure", stats["heures"]),
        ("Ventes moyennes par jour, année par année", stats["annees"]),
        ("Ventes de chaque mois complet, recalculées à partir des jours", stats["mensuel"]),
    ]


def construire(dossier_data, dossier_sortie):
    exports = charger(dossier_data)
    controles, mois = controler(exports)
    stats = statistiques(exports)
    s = saisons(exports, stats)
    jour = donnees_jour(exports)
    ecarts = mois[mois["Monthly en écart"] == "oui"].drop(columns=["Monthly en écart", "Jours du mois"])
    etat = [
        ("Quel export sert à quoi", etat_des_exports(exports)),
        ("Ce qui a été corrigé ou écarté", decisions()),
        ("Résultat de chaque contrôle", controles),
        ("Mois où l'export mensuel ne concorde pas avec les jours", ecarts),
    ]

    dossier_sortie.mkdir(parents=True, exist_ok=True)
    client = dossier_sortie / "Pharma_Analyse.xlsx"
    with pd.ExcelWriter(client, engine="openpyxl", datetime_format="DD/MM/YYYY") as xl:
        ws = {nom: xl.book.create_sheet(nom) for nom in ONGLETS if nom != "Données"}
        ecrire_synthese(ws["Synthèse"], exports, controles, mois, stats, s)
        ecrire_ventes(ws["Ventes"], jour)
        ecrire_calendrier(ws["Calendrier"], s)
        ecrire_tableaux(ws["Conseils"], "Conseils de commande par groupe",
                        f"Évolution sur un an : {s['periode']}.", [(None, s["conseils"])], largeur_max=70)
        ecrire_tableaux(ws["Questions"], "Questions fréquentes", None, [(None, faq(exports, mois))], largeur_max=95)
        ecrire_tableaux(ws["Médicaments"], "Ce que contient chaque groupe",
                        "Source : classification ATC de l'Organisation mondiale de la santé.", [(None, groupes_atc())])
        ecrire_tableaux(ws["Méthode"], "Méthode : comment chaque chiffre est obtenu", None, [(None, methode())],
                        largeur_max=75)
        ecrire_tableaux(ws["Statistiques"], "Statistiques détaillées",
                        "Quantités dans l'unité du logiciel de caisse. Colonnes : codes des groupes (voir Médicaments).",
                        blocs_statistiques(stats), largeur_max=38)
        ecrire_tableaux(ws["État des données"], "État des données", None, etat, largeur_max=95)
        jour.to_excel(xl, sheet_name="Données", index=False)
        entetes(xl.sheets["Données"], 1, list(jour.columns))
        xl.sheets["Données"].freeze_panes = "A2"
        for nom, (_, couleur, _) in ONGLETS.items():
            xl.sheets[nom].sheet_properties.tabColor = couleur

    traitement = dossier_sortie / "Pharma_Traitement.xlsx"
    with pd.ExcelWriter(traitement, engine="openpyxl") as xl:
        for nom, df in (("Controles", controles), ("Mois", mois), ("Exports", etat_des_exports(exports)),
                        ("Decisions", decisions())):
            df.to_excel(xl, sheet_name=nom, index=False)
            for j, col in enumerate(df.columns, start=1):
                largeur = max([len(str(col))] + [len(str(v)) for v in df[col]])
                xl.sheets[nom].column_dimensions[get_column_letter(j)].width = min(largeur + 2, 90)
            entetes(xl.sheets[nom], 1, list(df.columns))
    return client, traitement, controles


if __name__ == "__main__":
    p = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    p.add_argument("--data", type=Path, default=ROOT / "data", help="dossier contenant les quatre exports")
    p.add_argument("--out", type=Path, default=ROOT / "output", help="dossier où écrire les deux classeurs")
    args = p.parse_args()
    try:
        client, traitement, controles = construire(args.data, args.out)
    except FileNotFoundError as e:
        raise SystemExit(f"Export introuvable : {e.filename}. Les quatre fichiers Pharma_Ventes_*.csv "
                         f"doivent être dans {args.data}.")
    except PermissionError as e:
        raise SystemExit(f"Impossible d'écrire {e.filename} : fermez ce classeur dans Excel puis relancez.")
    print(f"Classeur client     : {client}")
    print(f"Classeur traitement : {traitement}")
    print(controles.to_string(index=False))
