"""Interface pour mettre à jour le classeur de la pharmacie et voir les résultats.

Lancement : double-clic sur mettre_a_jour.bat (ou : python interface.py).

Trois onglets :
- « Mise à jour » : importer les nouveaux exports CSV (quel que soit leur nom), puis mettre à jour le classeur
  (l'interface lance classeur.ipynb en arrière-plan) ;
- « Résultats » : les chiffres clés, la prévision du mois suivant et un graphique par groupe ;
- « Catégories » : les médicaments regroupés par catégorie (code ATC) et les nouveaux médicaments détectés.
"""

import os
import re
import shutil
import subprocess
import sys
import threading
from datetime import datetime
from pathlib import Path

import openpyxl
import pandas as pd

import categories as cat

DOSSIER = Path(__file__).resolve().parent          # le dossier rush-2
DATA = DOSSIER / "data"
CLASSEUR = DOSSIER / "excel" / "Pharma_Analyse.xlsx"
NOTEBOOK = DOSSIER / "classeur.ipynb"

GROUPES = cat.GROUPES_CONNUS
NOMS = {"M01AB": "Diclofénac", "M01AE": "Ibuprofène", "N02BA": "Aspirine", "N02BE": "Paracétamol",
        "N05B": "Anxiolytiques", "N05C": "Somnifères", "R03": "Asthme", "R06": "Antihistaminiques"}
NOMS_FICHIERS = {
    "Journalier": "Pharma_Ventes_Daily.csv",
    "Horaire": "Pharma_Ventes_Hourly.csv",
    "Hebdomadaire": "Pharma_Ventes_Weekly.csv",
    "Mensuel": "Pharma_Ventes_Monthly.csv",
}
MOTS_CLES = [
    ("Horaire", ["hour", "heure", "horaire"]),
    ("Hebdomadaire", ["week", "hebdo", "semaine"]),
    ("Mensuel", ["month", "mensuel", "mois"]),
    ("Journalier", ["daily", "day", "jour"]),
]
MOIS_FR = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
           "septembre", "octobre", "novembre", "décembre"]


# ======================================================================= la logique (sans fenêtre)

def lire_dates(colonne):
    """Lit la colonne des dates, quel que soit le format utilisé par le logiciel de caisse."""
    for fmt in ("%m/%d/%Y %H:%M", "%m/%d/%Y", "%Y-%m-%d", "%Y-%m-%d %H:%M:%S"):
        try:
            return pd.to_datetime(colonne, format=fmt)
        except (ValueError, TypeError):
            continue
    raise ValueError("format de date non reconnu")


def reconnaitre(chemin):
    """Renvoie (type d'export, nombre de lignes, première date, dernière date, nouveaux codes)."""
    table = pd.read_csv(chemin)
    manquants = [c for c in ["datum"] + GROUPES if c not in table.columns]
    if manquants:
        raise ValueError(f"ce n'est pas un export de ventes complet (colonnes absentes : {', '.join(manquants)})")
    dates = lire_dates(table["datum"])
    nom = Path(chemin).name.lower()
    type_export = next((t for t, mots in MOTS_CLES if any(m in nom for m in mots)), None)
    if type_export is None:                         # rien dans le nom : on regarde l'écart entre deux dates
        if (dates.dt.hour != 0).any():
            type_export = "Horaire"
        else:
            ecart = dates.sort_values().diff().dt.days.median()
            type_export = "Journalier" if ecart <= 1 else "Hebdomadaire" if ecart <= 7 else "Mensuel"
    return type_export, len(table), dates.min(), dates.max(), cat.nouveaux_codes(table.columns)


def importer(chemins):
    """Copie chaque CSV choisi dans data/ sous le nom attendu. Renvoie les lignes du compte rendu."""
    DATA.mkdir(exist_ok=True)
    compte_rendu, vus = [], {}
    for chemin in chemins:
        nom = Path(chemin).name
        try:
            type_export, n, debut, fin, nouveaux = reconnaitre(chemin)
        except Exception as erreur:
            compte_rendu.append(f"✗ {nom} : refusé, {erreur}")
            continue
        if type_export in vus:
            compte_rendu.append(f"✗ {nom} : refusé, un export {type_export.lower()} a déjà été choisi ({vus[type_export]})")
            continue
        vus[type_export] = nom
        cible = DATA / NOMS_FICHIERS[type_export]
        if Path(chemin).resolve() != cible.resolve():
            shutil.copyfile(chemin, cible)
        compte_rendu.append(f"✓ {type_export} : {n} lignes, du {debut:%d/%m/%Y} au {fin:%d/%m/%Y}  →  data/{cible.name}")
        if nouveaux:
            compte_rendu.append(f"   ★ Nouveaux médicaments dans cet export : {', '.join(nouveaux)} (voir l'onglet Catégories)")
    manquants = [t for t in ("Journalier", "Horaire") if t not in vus]
    if manquants and vus:
        compte_rendu.append(f"! Pas de nouvel export {' ni '.join(m.lower() for m in manquants)} : l'ancien est gardé.")
    return compte_rendu


def verifier():
    """Contrôle rapide : l'horaire additionné par jour doit donner le journalier."""
    jour = pd.read_csv(DATA / NOMS_FICHIERS["Journalier"])
    heure = pd.read_csv(DATA / NOMS_FICHIERS["Horaire"])
    jour.index = lire_dates(jour["datum"]).dt.normalize()
    heure_par_jour = heure[GROUPES].groupby(lire_dates(heure["datum"]).dt.normalize()).sum()
    communs = jour.index.intersection(heure_par_jour.index)
    ecarts = ((jour.loc[communs, GROUPES] - heure_par_jour.loc[communs, GROUPES]).abs() > 0.01).any(axis=1).sum()
    fin_jour, fin_heure = jour.index.max(), heure_par_jour.index.max()
    if fin_jour != fin_heure:
        return (f"! Contrôle : le journalier s'arrête le {fin_jour:%d/%m/%Y} et l'horaire le {fin_heure:%d/%m/%Y}. "
                "Importez les deux exports du même envoi.")
    if ecarts == 0:
        return f"✓ Contrôle : l'horaire et le journalier concordent sur {len(communs)} jours."
    return f"! Contrôle : {ecarts} jours où l'horaire et le journalier ne concordent pas. À signaler au client."


def excel_ouvert():
    """Excel crée un fichier ~$… à côté du classeur tant qu'il est ouvert."""
    return (CLASSEUR.parent / f"~${CLASSEUR.name}").exists()


def mettre_a_jour():
    """Lance classeur.ipynb sans ouvrir Jupyter. Renvoie (réussi ?, message)."""
    if not NOTEBOOK.exists():
        return False, "le fichier classeur.ipynb est introuvable dans le dossier rush-2 (vérifiez son nom exact)."
    if not (DOSSIER / "categories.py").exists():
        return False, "le fichier categories.py manque à côté de classeur.ipynb (dossier rush-2)."
    python = Path(sys.executable)
    if python.name.lower() == "pythonw.exe":                      # pythonw n'affiche pas les erreurs : on prend python.exe
        python = python.with_name("python.exe")
    commande = [str(python), "-m", "jupyter", "nbconvert", "--to", "notebook",
                "--execute", "--inplace", str(NOTEBOOK)]
    resultat = subprocess.run(commande, cwd=DOSSIER, capture_output=True, text=True, encoding="utf-8",
                              errors="replace", creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    if resultat.returncode != 0:
        sortie = re.sub(r"\x1b\[[0-9;]*m", "", (resultat.stderr or "") + "\n" + (resultat.stdout or "")).strip().splitlines()
        erreurs = [l.strip() for l in sortie if re.search(r"[A-Za-z]+(Error|Exception)\b", l) and "Equivalent" not in l]
        if erreurs:
            return False, erreurs[-1]
        return False, sortie[-1] if sortie else "erreur inconnue (aucun message)."
    return True, "✓ Classeur mis à jour. Les onglets Résultats et Catégories sont rafraîchis."


def charger_resultats():
    """Lit les ventes (CSV) et la prévision (classeur) pour les onglets Résultats et Catégories."""
    jour = pd.read_csv(DATA / NOMS_FICHIERS["Journalier"])
    jour["datum"] = lire_dates(jour["datum"])
    par_jour = jour.set_index("datum")[GROUPES]
    mois = par_jour.resample("MS").sum()
    nb = par_jour.resample("MS").size()
    mois = mois[nb == nb.index.days_in_month]

    wb = openpyxl.load_workbook(CLASSEUR, read_only=True, data_only=True)
    prev = wb["Prévision"]
    lignes = []
    for i, g in enumerate(GROUPES):
        lignes.append({
            "code": g, "nom": NOMS[g],
            "naive": prev.cell(16 + i, 7).value, "nous": prev.cell(16 + i, 8).value,
            "verdict": prev.cell(16 + i, 10).value,
            "prevision": prev.cell(27 + i, 3).value, "reel": prev.cell(27 + i, 4).value,
        })
    titre_prevision = prev["A25"].value
    wb.close()
    return {
        "jour": jour, "mois": mois, "prevision": lignes, "titre_prevision": titre_prevision,
        "dernier_jour": jour["datum"].max(), "mois_cible": mois.index.max() + pd.DateOffset(months=1),
        "categories": cat.ventes_par_categorie(jour), "codes": cat.tableau_codes(jour.columns),
        "nouveaux": cat.nouveaux_codes(jour.columns),
        "maj": datetime.fromtimestamp(CLASSEUR.stat().st_mtime),
    }


def ouvrir(chemin):
    if sys.platform.startswith("win"):
        os.startfile(chemin)
    else:
        subprocess.run(["open" if sys.platform == "darwin" else "xdg-open", str(chemin)])


# ======================================================================= la fenêtre

def lancer_fenetre():
    import tkinter as tk
    from tkinter import filedialog, messagebox, ttk

    from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
    from matplotlib.figure import Figure

    BLEU, BLEU_FONCE, BLEU_CLAIR = "#1F5F8B", "#174A6C", "#CFE3F5"
    FOND, BLANC, TEXTE, GRIS = "#F4F6F8", "#FFFFFF", "#1E2A36", "#5F6B76"
    VERT, JAUNE, ROUGE, ORANGE = "#C6EFCE", "#FFEB9C", "#FFC7CE", "#E07B00"
    POLICE = "Segoe UI" if sys.platform.startswith("win") else "DejaVu Sans"
    COULEURS_CAT = ["#1F5F8B", "#E07B00", "#2E8B57", "#8E44AD", "#C0392B", "#7F8C8D"]

    racine = tk.Tk()
    racine.title("Pharmacie – Mise à jour et résultats")
    racine.geometry("1180x800")
    racine.minsize(1000, 700)
    racine.configure(bg=FOND)

    style = ttk.Style()
    style.theme_use("clam")
    style.configure("TNotebook", background=FOND, borderwidth=0)
    style.configure("TNotebook.Tab", font=(POLICE, 11, "bold"), padding=(18, 8), background=BLEU_CLAIR, foreground=BLEU)
    style.map("TNotebook.Tab", background=[("selected", BLANC)], foreground=[("selected", BLEU_FONCE)])
    style.configure("Treeview", font=(POLICE, 10), rowheight=28, background=BLANC, fieldbackground=BLANC)
    style.configure("Treeview.Heading", font=(POLICE, 10, "bold"), background=BLEU_CLAIR, foreground=BLEU)
    style.configure("TCombobox", padding=4)
    style.configure("Horizontal.TProgressbar", troughcolor=BLEU_CLAIR, background=BLEU)

    # ---------- en-tête
    entete = tk.Frame(racine, bg=BLEU)
    entete.pack(fill="x")
    tk.Label(entete, text="Pharmacie · Suivi des ventes", bg=BLEU, fg=BLANC,
             font=(POLICE, 20, "bold")).pack(anchor="w", padx=24, pady=(14, 0))
    sous_titre = tk.Label(entete, text="", bg=BLEU, fg=BLEU_CLAIR, font=(POLICE, 11))
    sous_titre.pack(anchor="w", padx=24, pady=(0, 14))

    onglets = ttk.Notebook(racine)
    onglets.pack(fill="both", expand=True, padx=16, pady=12)

    def carte(parent, **kw):
        return tk.Frame(parent, bg=BLANC, highlightbackground="#D5DBE1", highlightthickness=1, **kw)

    def bouton(parent, texte, action, principal=True):
        b = tk.Button(parent, text=texte, command=action, font=(POLICE, 11, "bold"), cursor="hand2",
                      relief="flat", padx=18, pady=8,
                      bg=BLEU if principal else BLANC, fg=BLANC if principal else BLEU,
                      activebackground=BLEU_FONCE if principal else BLEU_CLAIR,
                      activeforeground=BLANC if principal else BLEU,
                      highlightbackground=BLEU, highlightthickness=0 if principal else 1)
        return b

    # ================================================================ onglet 1 : Mise à jour
    page_maj = tk.Frame(onglets, bg=FOND)
    onglets.add(page_maj, text="  1 · Mise à jour  ")

    def etape(numero, titre, texte):
        c = carte(page_maj)
        c.pack(fill="x", padx=8, pady=6)
        tk.Label(c, text=str(numero), bg=BLEU, fg=BLANC, font=(POLICE, 16, "bold"), width=3,
                 height=1).pack(side="left", fill="y")
        corps = tk.Frame(c, bg=BLANC)
        corps.pack(side="left", fill="both", expand=True, padx=16, pady=10)
        tk.Label(corps, text=titre, bg=BLANC, fg=TEXTE, font=(POLICE, 13, "bold")).pack(anchor="w")
        tk.Label(corps, text=texte, bg=BLANC, fg=GRIS, font=(POLICE, 10)).pack(anchor="w")
        zone = tk.Frame(c, bg=BLANC)
        zone.pack(side="right", padx=16)
        return zone

    zone1 = etape(1, "Importer les nouveaux exports",
                  "Choisissez les fichiers CSV reçus du client, quel que soit leur nom : ils sont reconnus et rangés dans data/.")
    zone2 = etape(2, "Mettre à jour le classeur",
                  "Fermez Excel, puis lancez la mise à jour (environ 30 secondes).")
    zone3 = etape(3, "Voir le résultat",
                  "Ouvrez le classeur Excel, ou consultez les onglets Résultats et Catégories ci-dessus.")

    progression = ttk.Progressbar(page_maj, mode="indeterminate", style="Horizontal.TProgressbar")
    journal = tk.Text(page_maj, height=10, font=("Consolas", 10), bg=BLANC, fg=TEXTE, relief="flat",
                      highlightbackground="#D5DBE1", highlightthickness=1, wrap="word", padx=10, pady=8)

    def ecrire(texte):
        journal.configure(state="normal")
        journal.insert("end", texte + "\n")
        journal.see("end")
        journal.configure(state="disabled")

    def action_importer():
        chemins = filedialog.askopenfilenames(title="Choisissez les nouveaux exports (CSV)",
                                              filetypes=[("Fichiers CSV", "*.csv")])
        if not chemins:
            return
        ecrire("— Import des exports —")
        for ligne in importer(chemins):
            ecrire(ligne)
        try:
            ecrire(verifier())
        except Exception as erreur:
            ecrire(f"! Contrôle impossible : {erreur}")
        ecrire("→ Étape suivante : Mettre à jour le classeur.")

    def action_mettre_a_jour():
        if excel_ouvert():
            messagebox.showwarning("Excel est ouvert", "Le classeur est ouvert dans Excel.\nFermez-le, puis cliquez à nouveau.")
            return
        bouton_maj.configure(state="disabled", text="Mise à jour en cours…")
        progression.pack(fill="x", padx=8, pady=(4, 0), before=journal)
        progression.start(12)
        ecrire("— Mise à jour du classeur —")

        def travail():
            try:
                reussi, message = mettre_a_jour()
            except Exception as erreur:
                reussi, message = False, str(erreur)

            def fin():
                progression.stop()
                progression.pack_forget()
                bouton_maj.configure(state="normal", text="Mettre à jour le classeur")
                if reussi:
                    ecrire(message)
                    rafraichir()
                    onglets.select(page_res)
                else:
                    ecrire(f"✗ La mise à jour a échoué : {message}\n"
                           "   Le plus souvent : un CSV au mauvais format. Ouvrez classeur.ipynb dans Jupyter pour le détail.")
            racine.after(0, fin)
        threading.Thread(target=travail, daemon=True).start()

    bouton(zone1, "Choisir les fichiers CSV…", action_importer).pack()
    bouton_maj = bouton(zone2, "Mettre à jour le classeur", action_mettre_a_jour)
    bouton_maj.pack()
    bouton(zone3, "Ouvrir le classeur Excel", lambda: ouvrir(CLASSEUR)).pack(side="left", padx=4)
    bouton(zone3, "Dossier data", lambda: ouvrir(DATA), principal=False).pack(side="left", padx=4)
    journal.pack(fill="both", expand=True, padx=8, pady=8)

    # ================================================================ onglet 2 : Résultats
    page_res = tk.Frame(onglets, bg=FOND)
    onglets.add(page_res, text="  2 · Résultats  ")

    bandeau = tk.Frame(page_res, bg=FOND)
    bandeau.pack(fill="x", padx=8, pady=(8, 4))
    chiffres = {}
    for cle, libelle in [("donnees", "Données jusqu'au"), ("cible", "Mois prévu"),
                         ("previsibles", "Groupes prévisibles"), ("para", "Paracétamol prévu")]:
        c = carte(bandeau)
        c.pack(side="left", fill="both", expand=True, padx=4)
        tk.Label(c, text=libelle, bg=BLANC, fg=GRIS, font=(POLICE, 10)).pack(anchor="w", padx=14, pady=(10, 0))
        chiffres[cle] = tk.Label(c, text="–", bg=BLANC, fg=BLEU, font=(POLICE, 20, "bold"))
        chiffres[cle].pack(anchor="w", padx=14, pady=(0, 10))

    milieu = tk.Frame(page_res, bg=FOND)
    milieu.pack(fill="both", expand=True, padx=8, pady=6)

    gauche = carte(milieu)
    gauche.pack(side="left", fill="both", padx=4)
    titre_tableau = tk.Label(gauche, text="Prévision du mois suivant", bg=BLANC, fg=TEXTE, font=(POLICE, 12, "bold"))
    titre_tableau.pack(anchor="w", padx=12, pady=(10, 4))
    colonnes = ("groupe", "prevision", "reel", "nous", "naive", "verdict")
    tableau = ttk.Treeview(gauche, columns=colonnes, show="headings", height=8)
    for col, titre, larg in zip(colonnes, ["Groupe", "Prévision", "Mois dernier", "Notre erreur", "Erreur naïve", "Prévisible ?"],
                                [130, 90, 115, 115, 115, 120]):
        tableau.heading(col, text=titre)
        tableau.column(col, width=larg, anchor="w" if col == "groupe" else "center")
    tableau.tag_configure("Oui", background=VERT)
    tableau.tag_configure("Gain marginal", background=JAUNE)
    tableau.tag_configure("Non", background=ROUGE)
    tableau.pack(fill="x", padx=12)
    tk.Label(gauche, text="Erreur = écart moyen entre la prévision et le réel, testé sur les derniers mois.\n"
                          "Vert : prévisible · Jaune : gain trop faible · Rouge : non prévisible.",
             bg=BLANC, fg=GRIS, font=(POLICE, 9), justify="left").pack(anchor="w", padx=12, pady=(6, 12))

    droite = carte(milieu)
    droite.pack(side="left", fill="both", expand=True, padx=4)
    haut_droite = tk.Frame(droite, bg=BLANC)
    haut_droite.pack(fill="x", padx=12, pady=(10, 0))
    tk.Label(haut_droite, text="Ventes mois par mois :", bg=BLANC, fg=TEXTE, font=(POLICE, 12, "bold")).pack(side="left")
    choix = tk.StringVar(value="Paracétamol")
    liste = ttk.Combobox(haut_droite, textvariable=choix, values=[NOMS[g] for g in GROUPES], state="readonly", width=18)
    liste.pack(side="left", padx=10)
    fig_res = Figure(figsize=(6, 3.6), dpi=100)
    toile_res = FigureCanvasTkAgg(fig_res, master=droite)
    toile_res.get_tk_widget().configure(highlightthickness=0, takefocus=0)
    toile_res.get_tk_widget().pack(fill="both", expand=True, padx=6, pady=6)

    # ================================================================ onglet 3 : Catégories
    page_cat = tk.Frame(onglets, bg=FOND)
    onglets.add(page_cat, text="  3 · Catégories  ")
    alerte = tk.Label(page_cat, text="", font=(POLICE, 11, "bold"), anchor="w", padx=14, pady=10)
    alerte.pack(fill="x", padx=12, pady=(10, 4))
    bas_cat = tk.Frame(page_cat, bg=FOND)
    bas_cat.pack(fill="both", expand=True, padx=8, pady=4)
    gauche_cat = carte(bas_cat)
    gauche_cat.pack(side="left", fill="both", padx=4)
    tk.Label(gauche_cat, text="Médicaments par catégorie (code ATC)", bg=BLANC, fg=TEXTE,
             font=(POLICE, 12, "bold")).pack(anchor="w", padx=12, pady=(10, 4))
    cols_cat = ("code", "sous", "cat", "nouveau")
    tableau_cat = ttk.Treeview(gauche_cat, columns=cols_cat, show="headings", height=12)
    for col, titre, larg in zip(cols_cat, ["Code", "Sous-groupe ATC", "Catégorie", "Nouveau ?"], [70, 300, 200, 85]):
        tableau_cat.heading(col, text=titre)
        tableau_cat.column(col, width=larg, anchor="center" if col in ("code", "nouveau") else "w")
    tableau_cat.tag_configure("nouveau", background=JAUNE)
    tableau_cat.pack(fill="both", expand=True, padx=12)
    tk.Label(gauche_cat, text="Un nouveau code dans l'export est rangé tout seul grâce au début de son code.\n"
                              "Les quantités de médicaments différents sont additionnées : à suivre dans le temps,\n"
                              "pas à comparer d'une catégorie à l'autre.",
             bg=BLANC, fg=GRIS, font=(POLICE, 9), justify="left").pack(anchor="w", padx=12, pady=(6, 12))
    droite_cat = carte(bas_cat)
    droite_cat.pack(side="left", fill="both", expand=True, padx=4)
    fig_cat = Figure(figsize=(6, 4), dpi=100)
    toile_cat = FigureCanvasTkAgg(fig_cat, master=droite_cat)
    toile_cat.get_tk_widget().configure(highlightthickness=0, takefocus=0)
    toile_cat.get_tk_widget().pack(fill="both", expand=True, padx=6, pady=6)

    # ================================================================ remplir avec les données
    donnees = {}

    def dessiner_groupe(*_):
        if not donnees:
            return
        code = next(g for g in GROUPES if NOMS[g] == choix.get())
        serie = donnees["mois"][code].iloc[-36:]
        fig_res.clear()
        ax = fig_res.add_subplot(111)
        ax.plot(serie.index, serie.values, color=BLEU, linewidth=2, label="Réel")
        ligne = next(l for l in donnees["prevision"] if l["code"] == code)
        if isinstance(ligne["prevision"], (int, float)):
            cible = donnees["mois_cible"]
            ax.plot([serie.index[-1], cible], [serie.iloc[-1], ligne["prevision"]], color=ORANGE, linestyle="--")
            ax.scatter([cible], [ligne["prevision"]], color=ORANGE, s=60, zorder=3,
                       label=f"Prévision : {ligne['prevision']}")
        else:
            ax.set_title("Pas de prévision fiable pour ce groupe : suivre le calendrier des saisons",
                         fontsize=9, color=GRIS)
        ax.grid(alpha=0.3)
        ax.legend(loc="upper left", fontsize=9)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        fig_res.autofmt_xdate()
        fig_res.tight_layout()
        toile_res.draw()

    liste.bind("<<ComboboxSelected>>", dessiner_groupe)

    def rafraichir():
        try:
            donnees.clear()
            donnees.update(charger_resultats())
        except Exception as erreur:
            sous_titre.configure(text=f"Résultats indisponibles : {erreur}. Lancez une mise à jour.")
            return
        d = donnees
        cible = d["mois_cible"]
        sous_titre.configure(text=f"Données jusqu'au {d['dernier_jour']:%d/%m/%Y}  ·  "
                                  f"classeur mis à jour le {d['maj']:%d/%m/%Y à %H:%M}")
        nb_oui = sum(1 for l in d["prevision"] if l["verdict"] == "Oui")
        para = next(l for l in d["prevision"] if l["code"] == "N02BE")["prevision"]
        chiffres["donnees"].configure(text=f"{d['dernier_jour']:%d/%m/%Y}")
        chiffres["cible"].configure(text=f"{MOIS_FR[cible.month - 1]} {cible.year}")
        chiffres["previsibles"].configure(text=f"{nb_oui} / 8")
        chiffres["para"].configure(text=f"{para}" if para is not None else "–")
        titre_tableau.configure(text=d["titre_prevision"].capitalize() if d["titre_prevision"] else "Prévision")

        tableau.delete(*tableau.get_children())
        pct = lambda x: f"{x * 100:.1f} %".replace(".", ",") if isinstance(x, (int, float)) else "–"
        for l in d["prevision"]:
            tableau.insert("", "end", values=(l["nom"], l["prevision"], l["reel"], pct(l["nous"]),
                                              pct(l["naive"]), l["verdict"]), tags=(l["verdict"],))
        dessiner_groupe()

        tableau_cat.delete(*tableau_cat.get_children())
        for _, r in d["codes"].iterrows():
            tableau_cat.insert("", "end", values=(r["Code"], r["Sous-groupe ATC"], r["Catégorie"], r["Nouveau ?"]),
                               tags=("nouveau",) if r["Nouveau ?"] == "Oui" else ())
        if d["nouveaux"]:
            alerte.configure(text=f"★ Nouveaux médicaments détectés : {', '.join(d['nouveaux'])}. "
                                  "Ils sont rangés dans leur catégorie ci-dessous.", bg=JAUNE, fg="#5C4500")
        else:
            alerte.configure(text="✓ Aucun nouveau médicament : l'export contient les mêmes 8 groupes qu'au départ.",
                             bg=VERT, fg="#1E5631")

        fig_cat.clear()
        ax = fig_cat.add_subplot(111)
        for i, nom in enumerate(d["categories"].columns):
            ax.plot(d["categories"].index, d["categories"][nom], linewidth=2,
                    color=COULEURS_CAT[i % len(COULEURS_CAT)], label=nom)
        ax.set_title("Ventes par catégorie, chaque mois complet", fontsize=11)
        ax.grid(alpha=0.3)
        ax.legend(fontsize=9)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        fig_cat.autofmt_xdate()
        fig_cat.tight_layout()
        toile_cat.draw()

    ecrire("Prêt. Pour un nouvel export : étape 1, puis étape 2.")
    rafraichir()
    racine.mainloop()


if __name__ == "__main__":
    lancer_fenetre()
